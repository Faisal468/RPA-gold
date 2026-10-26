
import configparser
import getpass
import queue
import re
import threading
from datetime import datetime, timedelta
from decimal import Decimal, ROUND_HALF_UP
from tkinter import messagebox

import customtkinter as ctk

import openpyxl
from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.by import By
from selenium.common.exceptions import TimeoutException, WebDriverException

import main as core
from main import BASE, LOG_FILE, LOG_SHEET, SetupError

CONFIG = BASE / "config.ini"


def load_settings():
    parser = configparser.ConfigParser(interpolation=None)
    if not parser.read(CONFIG, encoding="utf-8"):
        raise SetupError(f"Cannot find {CONFIG}")
    return parser.getint("settings", "difference_limit", fallback=5000)


def save_proof(driver, row_xpath, filename):
    """Screenshot the rate table (target row highlighted) into proof\\<dd Mon yyyy>\\<filename>."""
    row = driver.find_element(By.XPATH, row_xpath)
    driver.execute_script(
        "arguments[0].scrollIntoView({block:'center'});"
        "arguments[0].querySelectorAll('td,th').forEach(c => c.style.backgroundColor = '#FFF176');", row)
    table = row.find_element(By.XPATH, "./ancestor::table[1]")
    folder = BASE / "proof" / datetime.now().strftime("%d %b %Y")
    folder.mkdir(parents=True, exist_ok=True)
    if not table.screenshot(str(folder / filename)):
        raise SetupError(f"Could not save the proof screenshot {filename}")


def start_browser():
    """Hidden browser. Microsoft Edge first: every Windows PC has it and it is signed, so Windows
    Smart App Control allows it. The unsigned Chrome in the browser folder is only the fallback."""
    common = ["--headless=new", "--window-size=1400,1000"]
    try:
        options = webdriver.EdgeOptions()
        options.page_load_strategy = "eager"
        for arg in common:
            options.add_argument(arg)
        service = webdriver.EdgeService()
        service.creation_flags = 0x08000000  # no console window
        return webdriver.Edge(service=service, options=options)
    except WebDriverException:
        options = webdriver.ChromeOptions()
        options.binary_location = str(BASE / "browser" / "chrome-win64" / "chrome.exe")
        options.page_load_strategy = "eager"
        for arg in common:
            options.add_argument(arg)
        service = Service(executable_path=str(BASE / "browser" / "chromedriver-win64" / "chromedriver.exe"))
        service.creation_flags = 0x08000000
        return webdriver.Chrome(service=service, options=options)


def fetch_rates():
    """Open a hidden browser, read both sites, return (forex, pakbiz) as ints."""
    driver = start_browser()
    driver.set_page_load_timeout(60)
    try:
        tabs = []
        for i, url in enumerate(core.URLS):
            if i > 0:
                driver.switch_to.new_window("tab")
            tabs.append(driver.current_window_handle)
            try:
                driver.get(url)
            except TimeoutException:
                pass
        driver.switch_to.window(tabs[0])
        forex = core.read_forex(driver)
        save_proof(driver, "//td[normalize-space()='per 10 Gram']/ancestor::tr[1]", "Forex.png")
        driver.switch_to.window(tabs[1])
        pakbiz = core.read_pakbiz(driver)
        save_proof(driver, "//td[normalize-space()='24K']/ancestor::tr[1]", "PakBiz.png")
    finally:
        driver.quit()
    return int(forex.replace(",", "")), int(pakbiz.replace(",", "")) 


def weekly_average(today_avg):
    """Average of the daily averages for this Mon-Sun week, counting today's live average."""
    today = datetime.now().date()
    monday = today - timedelta(days=today.weekday())
    sunday = monday + timedelta(days=6)
    days = {}
    try:
        wb = openpyxl.load_workbook(LOG_FILE, read_only=True)
        ws = wb[LOG_SHEET] if LOG_SHEET in wb.sheetnames else wb.active
        for date, f, p in ((r[1], r[3], r[4]) for r in ws.iter_rows(min_row=2, max_col=5, values_only=True)):
            if isinstance(date, datetime) and monday <= date.date() <= sunday \
                    and isinstance(f, (int, float)) and isinstance(p, (int, float)):
                days[date.date()] = (f + p) / 2
        wb.close()
    except FileNotFoundError:
        pass
    days[today] = today_avg
    return sum(days.values()) / len(days)


SAME_RATE_DAYS = 6


def same_rate_alerts(forex, pakbiz):
    """Warnings for any source whose rate in the Excel log is unchanged for SAME_RATE_DAYS or more
    logged days in a row, counting today's live rates."""
    rates = {}  # date -> (forex, pakbiz)
    try:
        wb = openpyxl.load_workbook(LOG_FILE, read_only=True)
        ws = wb[LOG_SHEET] if LOG_SHEET in wb.sheetnames else wb.active
        for r in ws.iter_rows(min_row=2, max_col=5, values_only=True):
            if isinstance(r[1], datetime) and isinstance(r[3], (int, float)) and isinstance(r[4], (int, float)):
                rates[r[1].date()] = (r[3], r[4])
        wb.close()
    except FileNotFoundError:
        pass
    rates[datetime.now().date()] = (forex, pakbiz)
    history = [rates[d] for d in sorted(rates)]

    alerts = []
    for i, name in enumerate(("Forex.pk", "PakBiz")):
        values = [h[i] for h in history]
        run = 0
        for v in reversed(values):
            if v != values[-1]:
                break
            run += 1
        if run >= SAME_RATE_DAYS:
            alerts.append(f"{name} has shown the same rate ({values[-1]:,}) for {run} days in a row.")
    return alerts


GOLD, GOLD_HOVER = "#B8860B", "#9A6F08"
GREEN, RED = "#2E7D32", "#C62828"
FONT = "Segoe UI"

ctk.set_appearance_mode("light")
ctk.set_default_color_theme("blue")


class App(ctk.CTk):
    W, H = 420, 520

    def __init__(self):
        super().__init__()
        self.title("Gold Rate - 24K / 10 G")
        self.configure(fg_color="#F4F1EA")
        self.resizable(False, False)
        x, y = (self.winfo_screenwidth() - self.W) // 2, (self.winfo_screenheight() - self.H) // 3
        self.geometry(f"{self.W}x{self.H}+{x}+{y}")
        self.limit = load_settings()
        self.user = getpass.getuser()  # Windows username, saved as "Submitted by"
        self.results = queue.Queue()
        self.frame = None
        self.show_loading()

    def new_frame(self, title, subtitle=""):
        if self.frame:
            self.frame.destroy()
        self.frame = ctk.CTkFrame(self, fg_color="white", corner_radius=16)
        self.frame.pack(fill="both", expand=True, padx=20, pady=20)
        ctk.CTkLabel(self.frame, text="24K / 10 G GOLD RATE", font=(FONT, 12, "bold"),
                     text_color=GOLD).pack(pady=(24, 0))
        ctk.CTkLabel(self.frame, text=title, font=(FONT, 24, "bold"),
                     text_color="#222").pack()
        if subtitle:
            ctk.CTkLabel(self.frame, text=subtitle, font=(FONT, 12), text_color="#777").pack(pady=(2, 0))
        return self.frame

    def button(self, parent, text, command):
        return ctk.CTkButton(parent, text=text, command=command, height=42, corner_radius=10,
                             font=(FONT, 14, "bold"), fg_color=GOLD, hover_color=GOLD_HOVER)

    # ---- screen 1: fetch rates ----
    def show_loading(self):
        f = self.new_frame("Please wait", "Fetching today's rates...")
        bar = ctk.CTkProgressBar(f, mode="indeterminate", width=260, progress_color=GOLD)
        bar.pack(pady=(110, 14))
        bar.start()
        ctk.CTkLabel(f, text="This can take up to a minute.", font=(FONT, 12), text_color="#888").pack()
        threading.Thread(target=self.worker, daemon=True).start()
        self.after(300, self.poll)

    def worker(self):
        for attempt in (1, 2):  
            try:
                self.results.put(("ok", fetch_rates()))
                return
            except Exception as exc:
                text = str(exc) or repr(exc)
                if attempt == 2:
                    if any(k in text for k in ("ERR_INTERNET_DISCONNECTED", "ERR_NAME_NOT_RESOLVED",
                                               "ERR_CONNECTION", "ERR_PROXY", "ERR_TIMED_OUT")):
                        text = "No internet connection. Please check your internet and click Retry."
                    else:
                        text = text.splitlines()[0]
                    self.results.put(("error", text))

    def poll(self):
        try:
            kind, data = self.results.get_nowait()
        except queue.Empty:
            self.after(300, self.poll)
            return
        if kind == "ok":
            try:  # save the fetched rates straight away, as the original version did
                core.log_to_excel(str(data[0]), str(data[1]))
            except SetupError as exc:
                messagebox.showwarning("Not saved yet", f"{exc}\n\nIt will be saved when you click Submit.")
            self.show_rates(*data)
            alerts = same_rate_alerts(*data)
            if alerts:
                self.after(300, lambda: messagebox.showwarning(
                    "Same rate alert", "\n\n".join(alerts) + "\n\nPlease check that the website is updating."))
        elif messagebox.askretrycancel("Could not get rates", data):
            self.show_loading()
        else:
            self.destroy()

    # ---- screen 2: rates + submit ----
    def show_rates(self, forex, pakbiz):
        average = (Decimal(forex) + Decimal(pakbiz)) / 2
        average = int(average.quantize(Decimal(1), rounding=ROUND_HALF_UP))
        weekly = int(round(weekly_average(average)))
        diff = average - weekly  # negative when today is below the weekly average
        self.needs_manual = abs(diff) > self.limit
        self.data = dict(forex=forex, pakbiz=pakbiz, average=average, weekly=weekly, diff=diff)
        color = RED if self.needs_manual else GREEN

        f = self.new_frame("Today's Rates", datetime.now().strftime("%A, %d %B %Y"))
        card = ctk.CTkFrame(f, fg_color="#FAF7F0", corner_radius=12)
        card.pack(fill="x", padx=24, pady=(14, 10))
        rows = [("Forex rate", forex, "#222"), ("PakBiz rate", pakbiz, "#222"), ("Average", average, GOLD),
                ("Weekly Avg", weekly, "#222"), ("Difference", diff, color)]
        for i, (label, value, c) in enumerate(rows):
            card.grid_columnconfigure(0, weight=1)
            ctk.CTkLabel(card, text=label, font=(FONT, 13), text_color="#666").grid(
                row=i, column=0, sticky="w", padx=(16, 0), pady=5)
            shown = f"({-value:,})" if label == "Difference" and value < 0 else f"{value:,}"
            ctk.CTkLabel(card, text=shown, font=(FONT, 15, "bold"), text_color=c).grid(
                row=i, column=1, sticky="e", padx=16)

        note = (f"Difference is above {self.limit:,}. Enter your own rate."
                if self.needs_manual else f"Difference is within {self.limit:,}. Just click Submit.")
        ctk.CTkLabel(f, text=note, font=(FONT, 12, "bold"), text_color=color).pack(pady=(4, 6))
        self.manual = ctk.CTkEntry(f, placeholder_text="Your rate (optional)",
                                   height=42, width=290, corner_radius=10, justify="center")
        self.manual.pack(pady=6)
        self.button(f, "Submit", self.submit).pack(padx=45, pady=(10, 0), fill="x")
        if self.needs_manual:
            self.manual.focus()
            self.bind("<Return>", lambda e: self.submit())

    def submit(self):
        d = self.data
        text = self.manual.get().replace(",", "").strip()
        if text:  # box filled: use the user's own rate, otherwise the average
            if not re.fullmatch(r"\d+", text):
                messagebox.showerror("Rate needed", "Enter your own rate as a whole number.")
                return
            final, manual = int(text), "Yes"
        else:
            final, manual = d["average"], "No"
        try:
            core.log_to_excel(str(d["forex"]), str(d["pakbiz"]),
                              extra={"Final rate": final, "Manual rate": manual, "Submitted by": self.user})
        except SetupError as exc:
            messagebox.showerror("Could not save", str(exc))
            return
        messagebox.showinfo("Saved", f"Submitted. Final rate: {final:,}")
        self.destroy()


def self_test():
    import sys, traceback
    out = BASE / "test_output.txt"
    try:
        out.write_text(f"rates: {fetch_rates()}", encoding="utf-8")
    except BaseException:
        out.write_text(traceback.format_exc(), encoding="utf-8")


def main():
    import sys
    if "--test" in sys.argv:
        return self_test()
    try:
        App().mainloop()
    except SetupError as exc:
        messagebox.showerror("Setup problem", str(exc))


if __name__ == "__main__":
    main()
